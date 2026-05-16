#!/usr/bin/env python3
"""
SOP Ingestion Script — Cashless Everywhere

Reads the Cashless Everywhere SOP .xlsx and generates a SQL migration file
with idempotent INSERTs into:
  - hospital.panels                       (49 rows — TPAs + Insurers from Sheet 1)
  - hospital.preauth_form_templates       (42 rows — PDFs from Sheet 3)
  - hospital.mou_templates                (11 rows — MoU types from Sheet 4)
  - hospital.panel_default_attributes     (~490 rows — SOP defaults per panel)

Usage:
    python3 ingest.py <path-to-sop.xlsx> <path-to-output.sql>

Output is idempotent — re-running for an SOP refresh emits ON CONFLICT clauses.
"""

from __future__ import annotations

import re
import sys
import uuid
from pathlib import Path

import openpyxl

SOP_VERSION = "2026-05"

# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def code_from_name(name: str) -> str:
    """Mechanical name → CODE conversion. Idempotent, stable."""
    code = re.sub(r"[^a-zA-Z0-9]+", "_", name.strip())
    code = code.upper().strip("_")
    # Collapse multiple underscores
    code = re.sub(r"_+", "_", code)
    return code


def first_email(cell_value: str | None) -> str | None:
    """Extract the first valid email from a possibly-multi-line / annotated cell."""
    if not cell_value:
        return None
    text = str(cell_value)
    # Match standard email pattern
    matches = re.findall(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}", text)
    return matches[0] if matches else None


def all_emails(cell_value: str | None) -> list[str]:
    """Extract all valid emails from a cell."""
    if not cell_value:
        return []
    text = str(cell_value)
    return re.findall(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}", text)


def first_int(cell_value: str | None) -> int | None:
    """Extract first integer from a cell that may contain annotations."""
    if cell_value is None:
        return None
    if isinstance(cell_value, (int, float)):
        return int(cell_value)
    match = re.search(r"\d+", str(cell_value))
    return int(match.group()) if match else None


def sql_quote(value: str | None) -> str:
    """SQL string-escape (single-quote safe)."""
    if value is None:
        return "NULL"
    text = str(value).replace("'", "''")
    return f"'{text}'"


def sql_text_array(values: list[str]) -> str:
    """Format a Postgres TEXT[] literal."""
    if not values:
        return "'{}'::text[]"
    escaped = [v.replace("'", "''") for v in values]
    return "ARRAY[" + ", ".join(f"'{v}'" for v in escaped) + "]::text[]"


def sql_int(value: int | None) -> str:
    return str(value) if value is not None else "NULL"


# --------------------------------------------------------------------------
# Sheet parsers
# --------------------------------------------------------------------------

def parse_entities(sheet) -> list[dict]:
    """Parse Sheet 1 — 49 entity rows."""
    entities = []
    # Row 1 = headers, Row 2 = legend row, Rows 3+ = data
    for row_idx in range(3, sheet.max_row + 1):
        name = sheet.cell(row=row_idx, column=2).value
        if not name or not str(name).strip():
            continue

        entities.append({
            "row": row_idx,
            "name": str(name).strip(),
            "aliases_raw": sheet.cell(row=row_idx, column=3).value,
            "type": sheet.cell(row=row_idx, column=4).value,
            "email_cell": sheet.cell(row=row_idx, column=5).value,
            "phone_cell": sheet.cell(row=row_idx, column=6).value,
            "secondary_contact": sheet.cell(row=row_idx, column=7).value,
            "portal": sheet.cell(row=row_idx, column=8).value,
            "planned_hours": sheet.cell(row=row_idx, column=9).value,
            "emergency_hours": sheet.cell(row=row_idx, column=10).value,
            "form_name": sheet.cell(row=row_idx, column=11).value,
            "step_by_step": sheet.cell(row=row_idx, column=12).value,
            "onboarding_docs": sheet.cell(row=row_idx, column=13).value,
            "mou_type": sheet.cell(row=row_idx, column=14).value,
            "watch_outs": sheet.cell(row=row_idx, column=15).value,
            "network_size": sheet.cell(row=row_idx, column=16).value,
        })
    return entities


def parse_forms(sheet) -> list[dict]:
    """Parse Sheet 3 — 42 pre-auth forms."""
    forms = []
    for row_idx in range(2, sheet.max_row + 1):
        entity_name = sheet.cell(row=row_idx, column=2).value
        form_name = sheet.cell(row=row_idx, column=4).value
        if not entity_name or not form_name:
            continue

        forms.append({
            "row": row_idx,
            "entity_name": str(entity_name).strip(),
            "entity_type": sheet.cell(row=row_idx, column=3).value,
            "form_name": str(form_name).strip(),
            "url": sheet.cell(row=row_idx, column=5).value,
            "format": sheet.cell(row=row_idx, column=6).value,
            "key_differences": sheet.cell(row=row_idx, column=7).value,
            "notes": sheet.cell(row=row_idx, column=8).value,
        })
    return forms


def parse_mous(sheet) -> list[dict]:
    """Parse Sheet 4 — 11 MoU types."""
    mous = []
    for row_idx in range(2, sheet.max_row + 1):
        applies_to = sheet.cell(row=row_idx, column=2).value
        mou_type = sheet.cell(row=row_idx, column=3).value
        if not applies_to or not mou_type:
            continue

        mous.append({
            "row": row_idx,
            "applies_to": str(applies_to).strip(),
            "mou_type": str(mou_type).strip(),
            "validity": sheet.cell(row=row_idx, column=4).value,
            "key_clauses": sheet.cell(row=row_idx, column=5).value,
            "rate_terms": sheet.cell(row=row_idx, column=6).value,
            "audit_rights": sheet.cell(row=row_idx, column=7).value,
            "submission_deadline": sheet.cell(row=row_idx, column=8).value,
            "payment_timeline": sheet.cell(row=row_idx, column=9).value,
            "source_template": sheet.cell(row=row_idx, column=10).value,
            "risk_level": sheet.cell(row=row_idx, column=11).value,
        })
    return mous


# --------------------------------------------------------------------------
# Code generation
# --------------------------------------------------------------------------

def panel_type_for(entity_type: str | None) -> str:
    """Map SOP 'Entity Type' to panel_type CHECK enum."""
    if not entity_type:
        return "insurance"
    et = str(entity_type).strip().upper()
    if "TPA" in et:
        return "tpa"
    if "INSURER" in et or "INSURANCE" in et:
        return "insurance"
    if "PSU" in et or "GOVT" in et or "GOVERNMENT" in et:
        return "government"
    return "insurance"


def normalize_risk_level(raw: str | None) -> str | None:
    """Extract just the risk level word from the SOP cell."""
    if not raw:
        return None
    text = str(raw).upper()
    for marker in ("LOW-MEDIUM", "HIGH", "MEDIUM", "LOW"):
        if marker in text:
            return marker
    return None


def extract_first_line(text: str | None) -> str:
    """For one-line display contexts, the first non-empty line."""
    if not text:
        return ""
    for line in str(text).splitlines():
        stripped = line.strip()
        if stripped:
            return stripped
    return ""


def find_form_for_entity(forms: list[dict], entity_name: str) -> dict | None:
    """Match an entity row to a form row by entity name."""
    for f in forms:
        if f["entity_name"].lower() == entity_name.lower():
            return f
    return None


def find_mou_for_entity(mous: list[dict], entity_name: str) -> dict | None:
    """Match an entity row to a MoU row by applies_to."""
    for m in mous:
        if m["applies_to"].lower() == entity_name.lower():
            return m
        if entity_name.lower() in m["applies_to"].lower():
            return m
    return None


# --------------------------------------------------------------------------
# SQL emission
# --------------------------------------------------------------------------

def emit_header(out, sop_path: str):
    out.write(f"""-- Migration 016: Seed Cashless Everywhere SOP data
-- Date: 2026-05-16
-- Source SOP: {sop_path}
-- SOP Version: {SOP_VERSION}
-- Description:
--   Generated by Backend/scripts/sop-ingest/ingest.py
--   Seeds:
--     * panels (49 SOP entities)
--     * preauth_form_templates (42 forms)
--     * mou_templates (11 MoU types)
--     * panel_default_attributes (~490 default values per panel)
--
--   Idempotent: ON CONFLICT clauses protect re-runs.
--   To refresh after SOP updates: re-run the ingest script, review the diff,
--   create migration 017 with the diffs.

BEGIN;

""")


def emit_panels(out, entities: list[dict]):
    out.write("-- ============================================================\n")
    out.write(f"-- panels — {len(entities)} entities from SOP Sheet 1\n")
    out.write("-- ============================================================\n\n")
    for e in entities:
        code = code_from_name(e["name"])
        ptype = panel_type_for(e["type"])
        out.write(
            f"INSERT INTO hospital.panels (id, name, code, panel_type, is_system_panel)\n"
            f"VALUES (gen_random_uuid(), {sql_quote(e['name'])}, {sql_quote(code)}, "
            f"{sql_quote(ptype)}, FALSE)\n"
            f"ON CONFLICT (code) DO NOTHING;\n\n"
        )


def emit_forms(out, forms: list[dict]):
    out.write("-- ============================================================\n")
    out.write(f"-- preauth_form_templates — {len(forms)} forms from SOP Sheet 3\n")
    out.write("-- ============================================================\n")
    out.write("-- s3_key is NULL initially; downloaded + populated by\n")
    out.write("-- Backend/scripts/sop-ingest/download-forms.py in a separate step.\n\n")
    seen_codes = set()
    for f in forms:
        # Derive a stable code from the entity name + 'PREAUTH' suffix
        base_code = code_from_name(f["entity_name"])
        code = f"{base_code}_PREAUTH"
        # Disambiguate if duplicate (rare, e.g. multi-form entities)
        suffix = 2
        while code in seen_codes:
            code = f"{base_code}_PREAUTH_{suffix}"
            suffix += 1
        seen_codes.add(code)

        out.write(
            f"INSERT INTO hospital.preauth_form_templates "
            f"(id, name, code, source_url, is_active, notes)\n"
            f"VALUES (gen_random_uuid(), {sql_quote(f['form_name'])}, {sql_quote(code)}, "
            f"{sql_quote(f['url'])}, TRUE, {sql_quote(f['key_differences'])})\n"
            f"ON CONFLICT (code) DO NOTHING;\n\n"
        )


def emit_mous(out, mous: list[dict]):
    out.write("-- ============================================================\n")
    out.write(f"-- mou_templates — {len(mous)} MoU types from SOP Sheet 4\n")
    out.write("-- ============================================================\n\n")
    seen = set()
    for m in mous:
        base_code = code_from_name(m["applies_to"])
        code = f"{base_code}_MOU"
        suffix = 2
        while code in seen:
            code = f"{base_code}_MOU_{suffix}"
            suffix += 1
        seen.add(code)

        risk = normalize_risk_level(m["risk_level"])
        risk_notes = extract_first_line(m["risk_level"])

        out.write(
            f"INSERT INTO hospital.mou_templates\n"
            f"  (id, name, code, applies_to_label, validity_text, key_clauses,\n"
            f"   rate_terms, audit_rights, submission_deadline, payment_timeline,\n"
            f"   source_template_url, risk_level, risk_notes, is_active)\n"
            f"VALUES (gen_random_uuid(),\n"
            f"  {sql_quote(m['mou_type'])},\n"
            f"  {sql_quote(code)},\n"
            f"  {sql_quote(m['applies_to'])},\n"
            f"  {sql_quote(m['validity'])},\n"
            f"  {sql_quote(m['key_clauses'])},\n"
            f"  {sql_quote(m['rate_terms'])},\n"
            f"  {sql_quote(m['audit_rights'])},\n"
            f"  {sql_quote(m['submission_deadline'])},\n"
            f"  {sql_quote(m['payment_timeline'])},\n"
            f"  {sql_quote(m['source_template'])},\n"
            f"  {sql_quote(risk)},\n"
            f"  {sql_quote(risk_notes)},\n"
            f"  TRUE)\n"
            f"ON CONFLICT (code) DO NOTHING;\n\n"
        )


def emit_default_attributes(out, entities: list[dict], forms: list[dict], mous: list[dict]):
    """Emit panel_default_attributes rows by joining panels to their SOP config."""
    out.write("-- ============================================================\n")
    out.write("-- panel_default_attributes — SOP defaults per panel\n")
    out.write("-- ============================================================\n")
    out.write("-- Joins by panel code + attribute definition key. Idempotent.\n\n")

    for e in entities:
        panel_code = code_from_name(e["name"])
        emails = all_emails(e["email_cell"])
        primary_email = emails[0] if emails else None
        cc_emails = emails[1:] if len(emails) > 1 else []
        planned_hours = first_int(e["planned_hours"])
        emergency_hours = first_int(e["emergency_hours"])

        # Determine claim_submission_method based on whether we have an email
        if primary_email:
            submission_method = "email"
        elif e["portal"] and "portal" in str(e["portal"]).lower():
            submission_method = "portal"
        else:
            submission_method = "email"  # default; hospital may override

        # Form + MoU references (matched by name)
        matched_form = find_form_for_entity(forms, e["name"])
        matched_mou = find_mou_for_entity(mous, e["name"])
        form_code = f"{code_from_name(matched_form['entity_name'])}_PREAUTH" if matched_form else None
        mou_code = f"{code_from_name(matched_mou['applies_to'])}_MOU" if matched_mou else None

        out.write(f"-- {e['name']} ({panel_code})\n")

        # claim_submission_method
        _emit_text_default(out, panel_code, "claim_submission_method", submission_method)

        # claim_submission_email (existing single-email attribute)
        if primary_email:
            _emit_text_default(out, panel_code, "claim_submission_email", primary_email)

        # cashless_email_to_list (new multi-recipient attribute)
        if emails:
            _emit_text_default(out, panel_code, "cashless_email_to_list", ", ".join(emails))

        # cashless_email_cc_list
        if cc_emails:
            _emit_text_default(out, panel_code, "cashless_email_cc_list", ", ".join(cc_emails))

        # Portal URL
        if e["portal"]:
            _emit_text_default(out, panel_code, "portal_url", str(e["portal"]).strip())

        # Pre-auth deadlines (hours; new attributes)
        if planned_hours is not None:
            _emit_text_default(out, panel_code, "preauth_deadline_planned_hours", str(planned_hours))
        if emergency_hours is not None:
            _emit_text_default(out, panel_code, "preauth_deadline_emergency_hours", str(emergency_hours))

        # Default subject template
        subject_default = "Pre-Auth | {hospital_name} | {patient_name} | Policy {policy_no}"
        _emit_text_default(out, panel_code, "cashless_subject_template", subject_default)

        # SOP knowledge fields
        if e["secondary_contact"]:
            _emit_text_default(out, panel_code, "sop_secondary_contacts", str(e["secondary_contact"]).strip())
        if e["step_by_step"]:
            _emit_text_default(out, panel_code, "sop_step_by_step", str(e["step_by_step"]).strip())
        if e["watch_outs"]:
            _emit_text_default(out, panel_code, "sop_watch_outs", str(e["watch_outs"]).strip())

        # Form + MoU choices — only emit when matched
        if form_code:
            # We store the form template's CODE here as a stable identifier.
            # The application resolves it to the actual UUID at runtime.
            _emit_text_default(out, panel_code, "preauth_form_template_id", form_code)
        if mou_code:
            _emit_text_default(out, panel_code, "mou_template_id", mou_code)

        out.write("\n")


def _emit_text_default(out, panel_code: str, attribute_key: str, text_value: str):
    """Emit one panel_default_attributes row for a text-valued attribute."""
    out.write(
        f"INSERT INTO hospital.panel_default_attributes\n"
        f"  (panel_id, panel_attribute_definition_id, default_value_text, source, sop_version)\n"
        f"SELECT p.id, pad.id, {sql_quote(text_value)}, 'sop_seed', {sql_quote(SOP_VERSION)}\n"
        f"FROM hospital.panels p, hospital.panel_attribute_definitions pad\n"
        f"WHERE p.code = {sql_quote(panel_code)} AND pad.key = {sql_quote(attribute_key)}\n"
        f"ON CONFLICT (panel_id, panel_attribute_definition_id) DO NOTHING;\n"
    )


def emit_footer(out):
    out.write("\nCOMMIT;\n")


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def main():
    if len(sys.argv) != 3:
        print("Usage: ingest.py <sop.xlsx> <output.sql>", file=sys.stderr)
        sys.exit(1)

    sop_path = sys.argv[1]
    out_path = sys.argv[2]

    wb = openpyxl.load_workbook(sop_path, data_only=True)
    entities = parse_entities(wb["1. Entity Lookup & Routing"])
    forms = parse_forms(wb["3. Pre-Auth Form Downloads (42)"])
    mous = parse_mous(wb["4. MoU Variations (11 types)"])

    print(f"Parsed: {len(entities)} entities, {len(forms)} forms, {len(mous)} MoUs", file=sys.stderr)

    with open(out_path, "w") as out:
        emit_header(out, sop_path)
        emit_panels(out, entities)
        emit_forms(out, forms)
        emit_mous(out, mous)
        emit_default_attributes(out, entities, forms, mous)
        emit_footer(out)

    print(f"Wrote {out_path}", file=sys.stderr)


if __name__ == "__main__":
    main()
